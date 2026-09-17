from __future__ import annotations

import json
import math
import uuid
from datetime import timedelta
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

from django.db import IntegrityError, transaction
from django.db.models import Max, Q
from django.utils import timezone

from .models import RobotProfile, RobotSchedule, Shelf, Warehouse, WarehouseOrder, WorkPoint
from .schedule_services import (
    ACTIVE_SCHEDULE_STATUSES,
    PRIORITY_RANK,
    _flow_kind,
    _shelf_for_workpoint,
    create_order,
    create_schedule,
    reserve_inventory_for_order,
    validate_flow_route,
)

IMPORT_EXTENSIONS = {'.json', '.xlsx'}
ACTIVE_ORDER_STATUSES = ('NEW', 'PLANNED', 'QUEUED', 'RUNNING')

# Spreadsheet/JSON aliases intentionally accept both business-friendly and API-like names.
ALIASES: dict[str, tuple[str, ...]] = {
    'order_no': ('order_no', 'order', 'order_id', 'order_number', 'ma_don', 'ma_don_hang'),
    'external_ref': ('external_ref', 'reference', 'ref', 'po', 'asn', 'so', 'customer_ref'),
    'item_code': ('item_code', 'sku', 'product_code', 'goods_code', 'ma_hang', 'ma_san_pham'),
    'item_name': ('item_name', 'product_name', 'goods_name', 'ten_hang', 'ten_san_pham'),
    'quantity': ('quantity', 'qty', 'so_luong'),
    'priority': ('priority', 'muc_uu_tien'),
    'payload_weight_kg': ('payload_weight_kg', 'weight_kg', 'weight', 'khoi_luong_kg'),
    'due_at': ('due_at', 'due_time', 'deadline', 'required_at'),
    'notes': ('notes', 'note', 'ghi_chu'),
    'shelf_code': (
        'shelf_code', 'shelf', 'rack_id', 'rack_code', 'location', 'location_code',
        'storage_location', 'current_location', 'vi_tri', 'vi_tri_hang',
    ),
    'dock_code': ('dock_code', 'dock', 'door', 'gate', 'dock_id'),
    'source_code': ('source_code', 'source', 'from', 'from_location'),
    'destination_code': ('destination_code', 'destination', 'to', 'to_location'),
    'type': ('type', 'flow', 'order_type'),
    'load_units': ('load_units', 'load_unit', 'units'),
}


def _clean_key(value: Any) -> str:
    return str(value or '').strip().lower().replace(' ', '_').replace('-', '_')


def _clean_value(value: Any) -> Any:
    if value is None:
        return ''
    if isinstance(value, str):
        return value.strip()
    return value


def _normalized_row(raw: dict[str, Any]) -> dict[str, Any]:
    normalized = {_clean_key(k): _clean_value(v) for k, v in raw.items() if str(k or '').strip()}
    out: dict[str, Any] = {'_raw': normalized}
    for canonical, aliases in ALIASES.items():
        for alias in aliases:
            key = _clean_key(alias)
            if key in normalized and normalized[key] not in ('', None):
                out[canonical] = normalized[key]
                break
    return out


def _decode_json(uploaded_file, flow: str) -> list[dict[str, Any]]:
    try:
        payload = json.loads(uploaded_file.read().decode('utf-8-sig'))
    except UnicodeDecodeError as exc:
        raise ValueError('JSON file must be UTF-8 encoded') from exc
    except json.JSONDecodeError as exc:
        raise ValueError(f'invalid JSON: {exc.msg}') from exc

    if isinstance(payload, list):
        rows = payload
    elif isinstance(payload, dict):
        candidates = [
            payload.get('orders'),
            payload.get(flow.lower()),
            payload.get(f'{flow.lower()}_orders'),
            payload.get('data'),
        ]
        rows = next((x for x in candidates if isinstance(x, list)), None)
        if rows is None:
            # A single order object is useful for quick tests and is unambiguous.
            rows = [payload]
    else:
        raise ValueError('JSON root must be an order object, an array, or {"orders": [...]}')

    output = []
    for index, row in enumerate(rows, start=1):
        if not isinstance(row, dict):
            raise ValueError(f'JSON row {index} must be an object')
        output.append(_normalized_row(row))
    return output


def _decode_xlsx(uploaded_file) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError('Excel import requires openpyxl; run pip install -r requirements.txt') from exc

    try:
        content = uploaded_file.read()
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f'cannot read Excel file: {exc}') from exc

    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header_values = next(rows)
    except StopIteration:
        return []
    headers = [_clean_key(v) for v in header_values]
    if not any(headers):
        raise ValueError('Excel first row must contain column headers')

    output: list[dict[str, Any]] = []
    for values in rows:
        if not any(v not in (None, '') for v in values):
            continue
        raw = {headers[i]: values[i] for i in range(min(len(headers), len(values))) if headers[i]}
        output.append(_normalized_row(raw))
    return output


def parse_order_file(uploaded_file, flow: str) -> list[dict[str, Any]]:
    flow = str(flow or '').upper()
    if flow not in ('INBOUND', 'OUTBOUND'):
        raise ValueError('flow must be INBOUND or OUTBOUND')
    name = str(getattr(uploaded_file, 'name', '') or '')
    ext = Path(name).suffix.lower()
    if ext not in IMPORT_EXTENSIONS:
        raise ValueError(f'{name or "file"}: only .xlsx and .json are supported')
    if getattr(uploaded_file, 'size', 0) and int(uploaded_file.size) > 10 * 1024 * 1024:
        raise ValueError(f'{name}: file is larger than 10 MB')
    if ext == '.json':
        rows = _decode_json(uploaded_file, flow)
    else:
        rows = _decode_xlsx(uploaded_file)
    if not rows:
        raise ValueError(f'{name}: no order rows found')
    if len(rows) > 2000:
        raise ValueError(f'{name}: maximum 2000 rows per file')
    return rows


def _resolve_workpoint(warehouse: Warehouse, token: Any, expected_kind: str) -> WorkPoint | None:
    value = str(token or '').strip()
    if not value:
        return None
    q = WorkPoint.objects.filter(warehouse=warehouse, enabled=True).filter(
        Q(code__iexact=value) | Q(resource_id__iexact=value) | Q(name__iexact=value)
    ).order_by('id')
    for point in q:
        if _flow_kind(point) == expected_kind:
            return point
    return None


def _default_dock(warehouse: Warehouse, flow: str) -> WorkPoint:
    for point in WorkPoint.objects.filter(warehouse=warehouse, enabled=True).order_by('code'):
        if _flow_kind(point) == flow:
            return point
    raise ValueError(f'no enabled {flow} dock is configured')


def _physical_shelf(point: WorkPoint) -> Shelf | None:
    return _shelf_for_workpoint(point)


def _pending_inbound_for(point: WorkPoint) -> int:
    return WarehouseOrder.objects.filter(
        warehouse=point.warehouse,
        type='INBOUND',
        destination__resource_id=point.resource_id,
        status__in=ACTIVE_ORDER_STATUSES,
    ).count()


def _auto_inbound_shelf(warehouse: Warehouse, dock: WorkPoint) -> WorkPoint:
    unique: dict[str, WorkPoint] = {}
    for point in WorkPoint.objects.filter(warehouse=warehouse, enabled=True).order_by('code'):
        if _flow_kind(point) != 'SHELF':
            continue
        key = str(point.resource_id or point.code)
        unique.setdefault(key, point)

    candidates: list[tuple[int, float, str, WorkPoint]] = []
    for point in unique.values():
        shelf = _physical_shelf(point)
        if shelf is None or shelf.status in ('DISABLED', 'MAINTENANCE'):
            continue
        reserved = _pending_inbound_for(point)
        effective = int(shelf.current_load) + reserved
        if effective >= 8:
            continue
        distance = math.hypot(float(dock.x) - float(point.x), float(dock.y) - float(point.y))
        candidates.append((effective, distance, str(shelf.code), point))

    if not candidates:
        raise ValueError('no shelf has a free order slot for inbound')
    # Balance occupancy first, then choose the nearest candidate. This keeps the
    # initial 0/1/2/3 demo distribution readable instead of filling one rack to 8.
    candidates.sort(key=lambda x: (x[0], x[1], x[2]))
    return candidates[0][3]


def _pick_flow_locations(warehouse: Warehouse, flow: str, row: dict[str, Any]) -> tuple[WorkPoint, WorkPoint]:
    row_type = str(row.get('type') or '').upper()
    if row_type and row_type != flow:
        raise ValueError(f'row type {row_type} does not match uploaded {flow} file')

    if flow == 'INBOUND':
        dock_token = row.get('dock_code') or row.get('source_code')
        shelf_token = row.get('shelf_code') or row.get('destination_code')
        dock = _resolve_workpoint(warehouse, dock_token, 'INBOUND') if dock_token else _default_dock(warehouse, 'INBOUND')
        if dock is None:
            raise ValueError(f'unknown inbound dock {dock_token}')
        if shelf_token:
            shelf = _resolve_workpoint(warehouse, shelf_token, 'SHELF')
            if shelf is None:
                raise ValueError(f'unknown destination shelf/location {shelf_token}')
        else:
            shelf = _auto_inbound_shelf(warehouse, dock)
        return dock, shelf

    shelf_token = row.get('shelf_code') or row.get('source_code')
    dock_token = row.get('dock_code') or row.get('destination_code')
    if not shelf_token:
        raise ValueError('OUTBOUND row requires shelf_code/location so the robot knows where the goods are')
    shelf = _resolve_workpoint(warehouse, shelf_token, 'SHELF')
    if shelf is None:
        raise ValueError(f'unknown source shelf/location {shelf_token}')
    dock = _resolve_workpoint(warehouse, dock_token, 'OUTBOUND') if dock_token else _default_dock(warehouse, 'OUTBOUND')
    if dock is None:
        raise ValueError(f'unknown outbound dock {dock_token}')
    return shelf, dock


def _priority(value: Any) -> str:
    v = str(value or 'NORMAL').upper().strip()
    if v not in dict(WarehouseOrder.PRIORITIES):
        raise ValueError(f'invalid priority {v}')
    return v


def _positive_int(value: Any, default: int = 1) -> int:
    if value in ('', None):
        return default
    try:
        result = int(float(value))
    except (TypeError, ValueError) as exc:
        raise ValueError(f'invalid integer value {value}') from exc
    return max(1, result)


def _nonnegative_float(value: Any, default: float = 0.0) -> float:
    if value in ('', None):
        return default
    try:
        return max(0.0, float(value))
    except (TypeError, ValueError) as exc:
        raise ValueError(f'invalid numeric value {value}') from exc


def _row_order_data(
    *, warehouse: Warehouse, flow: str, row: dict[str, Any], source: WorkPoint, destination: WorkPoint,
    batch_id: str, filename: str, row_number: int,
) -> dict[str, Any]:
    raw = dict(row.get('_raw') or {})
    item_code = str(row.get('item_code') or '').strip()
    item_name = str(row.get('item_name') or '').strip()
    shelf_point = destination if flow == 'INBOUND' else source
    shelf = _physical_shelf(shelf_point)
    dock_point = source if flow == 'INBOUND' else destination
    metadata = {
        'created_from': 'order_file_import',
        'import_batch_id': batch_id,
        'import_filename': filename,
        'import_row': row_number,
        'item_code': item_code,
        'item_name': item_name,
        'shelf_code': (shelf.layout_rack_id or shelf.code) if shelf else str(shelf_point.resource_id or shelf_point.code),
        'dock_code': dock_point.code,
        'shelf_slot_units': 1,
        'import_payload': raw,
    }
    order_no = str(row.get('order_no') or '').strip() or f'{flow[:3]}-{timezone.now().strftime("%Y%m%d")}-{uuid.uuid4().hex[:7].upper()}'
    load_units = _positive_int(row.get('load_units'), 1)
    if load_units > 4:
        raise ValueError('load_units must be between 1 and 4')
    return {
        'warehouse_id': warehouse.id,
        'order_no': order_no,
        'external_ref': str(row.get('external_ref') or '').strip(),
        'type': flow,
        'priority': _priority(row.get('priority')),
        'source_id': source.id,
        'destination_id': destination.id,
        'quantity': _positive_int(row.get('quantity'), 1),
        'load_units': load_units,
        'payload_weight_kg': _nonnegative_float(row.get('payload_weight_kg'), 0.0),
        'due_at': row.get('due_at') or None,
        'notes': str(row.get('notes') or '').strip(),
        'metadata': metadata,
    }


def _next_robot_free_time(profile: RobotProfile, earliest):
    latest = RobotSchedule.objects.filter(
        robot=profile,
        status__in=ACTIVE_SCHEDULE_STATUSES,
        planned_end__isnull=False,
    ).aggregate(value=Max('planned_end'))['value']
    if latest and latest >= earliest:
        return latest + timedelta(seconds=1)
    return earliest


def _robot_runtime_eligible(
    profile: RobotProfile, order: WarehouseOrder, runtime_state: dict[str, Any], *, prefer_unloaded: bool = False
) -> tuple[bool, str | None]:
    state = runtime_state.get(profile.robot_id) or {}
    if not state:
        return False, 'robot is offline'
    status = str(state.get('status') or 'OFFLINE').upper()
    if status in ('OFFLINE', 'ERROR'):
        return False, f'robot is {status.lower()}'
    battery = float(state.get('battery') or 0.0)
    if battery < float(profile.min_dispatch_battery):
        return False, f'battery {battery:.0f}% below {profile.min_dispatch_battery:.0f}%'
    if int(order.load_units) > int(profile.payload_capacity):
        return False, f'load {order.load_units} exceeds capacity {profile.payload_capacity}'
    if prefer_unloaded:
        carrying = float((state.get('load') or {}).get('current') or 0.0)
        fsm = str(state.get('fsm') or '').upper()
        # "Prefer unloaded" is intentionally strict about delivery: a robot that
        # already carries goods or is transporting/delivering them is never selected.
        # A robot on its way to a pickup (TASK_ASSIGNED/NAVIGATING/PICKING) is allowed
        # and its new job is queued after its current planned work.
        if carrying > 0 or fsm in ('TRANSPORTING', 'DELIVERING'):
            return False, 'robot is currently carrying/delivering another order'
    return True, None


def _unloaded_dispatch_rank(state: dict[str, Any]) -> int:
    """Lower is better for the operator's unloaded-robot preference."""
    carrying = float((state.get('load') or {}).get('current') or 0.0)
    fsm = str(state.get('fsm') or 'OFFLINE').upper()
    if carrying > 0 or fsm in ('TRANSPORTING', 'DELIVERING'):
        return 99
    if fsm in ('IDLE', 'COMPLETED') and not state.get('current_task_id'):
        return 0
    if fsm in ('TASK_ASSIGNED', 'NAVIGATING', 'PICKING', 'OBSTACLE_DETECTED', 'REPLANNING'):
        return 1
    return 2


def auto_schedule_order(
    order: WarehouseOrder, runtime, user=None, *, earliest_start=None, prefer_unloaded_robot: bool = False
) -> RobotSchedule:
    """Create the earliest conflict-free AUTO schedule for one imported order.

    Future jobs are queued after each robot's last planned job. If a shared dock
    reservation conflicts, the start is shifted forward and retried. This makes
    bulk imports deterministic instead of repeatedly failing on the first robot.
    """
    if order.schedules.filter(status__in=ACTIVE_SCHEDULE_STATUSES).exists():
        return order.schedules.filter(status__in=ACTIVE_SCHEDULE_STATUSES).order_by('planned_start', 'id').first()

    base = earliest_start or timezone.now()
    profiles = list(RobotProfile.objects.filter(warehouse=order.warehouse, enabled=True).order_by('robot_id'))
    runtime_state = runtime.engine.state.get('robots') or {}
    eligible: list[tuple[int, Any, float, RobotProfile]] = []
    reasons: list[str] = []
    for profile in profiles:
        ok, reason = _robot_runtime_eligible(
            profile, order, runtime_state, prefer_unloaded=prefer_unloaded_robot
        )
        if not ok:
            reasons.append(f'{profile.robot_id}: {reason}')
            continue
        state = runtime_state.get(profile.robot_id) or {}
        free_at = _next_robot_free_time(profile, base)
        # If the simulator shows an untracked task, do not pretend the robot can
        # start immediately. eta_s is advisory, so add a small service buffer.
        if state.get('current_task_id') and not RobotSchedule.objects.filter(
            robot=profile, status__in=ACTIVE_SCHEDULE_STATUSES
        ).exists():
            eta = max(0.0, float(state.get('eta_s') or 0.0))
            free_at = max(free_at, base + timedelta(seconds=eta + 30.0))
        pos = state.get('position') or [0, 0, 0]
        distance_to_pickup = math.hypot(float(pos[0]) - float(order.source.x), float(pos[2]) - float(order.source.y))
        rank = _unloaded_dispatch_rank(state) if prefer_unloaded_robot else 0
        eligible.append((rank, free_at, distance_to_pickup, profile))
    if not eligible:
        prefix = 'no unloaded/non-delivering robot is eligible' if prefer_unloaded_robot else 'no eligible robot'
        raise ValueError(prefix + ': ' + ('; '.join(reasons) if reasons else 'no enabled robot profiles'))

    # With the preference enabled: truly idle/unloaded first, then robots going to
    # pickup, then other unloaded states. Within a class choose earliest + nearest.
    eligible.sort(key=lambda x: (x[0], x[1], x[2], x[3].robot_id))
    errors: list[str] = []
    for _rank, start, _distance, profile in eligible:
        attempt_start = start
        for _ in range(120):  # up to ~60 minutes in 30-second increments
            try:
                return create_schedule({
                    'order_id': order.id,
                    'planned_start': attempt_start.isoformat(),
                    'mode': 'AUTO',
                    'robot_id': profile.robot_id,
                }, runtime, user)
            except ValueError as exc:
                text = str(exc)
                if 'resource conflict' in text or 'conflicting schedule' in text or 'schedule conflict' in text:
                    attempt_start += timedelta(seconds=30)
                    continue
                errors.append(f'{profile.robot_id}: {text}')
                break
    raise ValueError('automatic scheduling failed' + (': ' + '; '.join(errors) if errors else ''))


def _sort_key(item: tuple[WarehouseOrder, dict[str, Any]]):
    order, context = item
    due = order.due_at or (timezone.now() + timedelta(days=3650))
    return (-PRIORITY_RANK.get(order.priority, 0), due, context.get('sequence', 0))


def import_and_auto_schedule(*, warehouse: Warehouse, runtime, user, files: Iterable[tuple[str, Any]]) -> dict[str, Any]:
    """Import INBOUND/OUTBOUND files, create orders, then schedule them automatically.

    `files` contains tuples like ('INBOUND', uploaded_file). Invalid files/rows are
    reported without aborting the other file. Re-uploading an existing order_no is
    idempotent: the row is marked duplicate and not recreated.
    """
    files = list(files)
    batch_id = f'IMP-{timezone.now().strftime("%Y%m%d-%H%M%S")}-{uuid.uuid4().hex[:6].upper()}'
    created: list[tuple[WarehouseOrder, dict[str, Any]]] = []
    results: list[dict[str, Any]] = []
    parsed_files: list[tuple[str, str, list[dict[str, Any]]]] = []
    sequence = 0

    # Parse everything before mutating the database. A malformed second file must
    # never strand orders created from the first file without a result payload.
    for flow, uploaded_file in files:
        flow = str(flow or '').upper()
        filename = str(getattr(uploaded_file, 'name', '') or f'{flow.lower()}.file')
        try:
            rows = parse_order_file(uploaded_file, flow)
            parsed_files.append((flow, filename, rows))
        except ValueError as exc:
            results.append({
                'flow': flow, 'filename': filename, 'row': 0,
                'status': 'FAILED', 'message': str(exc),
            })

    for flow, filename, rows in parsed_files:
        for row_number, row in enumerate(rows, start=2 if filename.lower().endswith('.xlsx') else 1):
            sequence += 1
            base_result = {'flow': flow, 'filename': filename, 'row': row_number, 'status': 'FAILED'}
            try:
                given_order_no = str(row.get('order_no') or '').strip()
                if given_order_no:
                    existing = WarehouseOrder.objects.filter(order_no=given_order_no).select_related('source', 'destination').first()
                    if existing:
                        base_result.update({
                            'status': 'DUPLICATE', 'order_id': existing.id, 'order_no': existing.order_no,
                            'message': 'order_no already exists; skipped',
                        })
                        results.append(base_result)
                        continue

                source, destination = _pick_flow_locations(warehouse, flow, row)
                # Validate before creating so the error points to the imported row.
                validate_flow_route(flow, source, destination)
                data = _row_order_data(
                    warehouse=warehouse, flow=flow, row=row, source=source, destination=destination,
                    batch_id=batch_id, filename=filename, row_number=row_number,
                )
                existing = WarehouseOrder.objects.filter(order_no=data['order_no']).select_related('source', 'destination').first()
                if existing:
                    base_result.update({
                        'status': 'DUPLICATE', 'order_id': existing.id, 'order_no': existing.order_no,
                        'message': 'order_no already exists; skipped',
                    })
                    results.append(base_result)
                    continue

                with transaction.atomic():
                    order = create_order(data, user)
                    if flow == 'OUTBOUND':
                        # Bind this outbound request to one exact physical item on
                        # the supplied shelf. When item_code is present it must match;
                        # otherwise the oldest unreserved item is selected.
                        reserve_inventory_for_order(
                            order, item_code=str(row.get('item_code') or '').strip()
                        )
                context = {'sequence': sequence, 'result': base_result}
                created.append((order, context))
                base_result.update({
                    'status': 'CREATED', 'order_id': order.id, 'order_no': order.order_no,
                    'source': order.source.code, 'destination': order.destination.code,
                    'shelf_code': order.metadata.get('shelf_code'), 'item_code': order.metadata.get('item_code'),
                })
                results.append(base_result)
            except (ValueError, KeyError, TypeError, IntegrityError) as exc:
                base_result['message'] = str(exc)
                results.append(base_result)

    # Higher priority / earlier due orders get the first robot time slots, regardless
    # of whether they came from the inbound or outbound file.
    created.sort(key=_sort_key)
    for order, context in created:
        result = context['result']
        try:
            schedule = auto_schedule_order(order, runtime, user, earliest_start=timezone.now())
            result.update({
                'status': 'SCHEDULED',
                'schedule_id': schedule.schedule_id,
                'robot_id': schedule.robot.robot_id,
                'planned_start': schedule.planned_start.isoformat(),
                'planned_end': schedule.planned_end.isoformat() if schedule.planned_end else None,
            })
        except ValueError as exc:
            # Keep the order NEW so the operator can schedule it later if robots are
            # offline. Import success and scheduling failure are deliberately separate.
            order.refresh_from_db()
            if order.status == 'PLANNED' and not order.schedules.filter(status__in=ACTIVE_SCHEDULE_STATUSES).exists():
                order.status = 'NEW'
                order.save(update_fields=['status', 'updated_at'])
            result.update({'status': 'UNSCHEDULED', 'message': str(exc)})

    summary = {
        'files': len(files),
        'rows': len(results),
        'created': sum(1 for r in results if r['status'] in ('CREATED', 'SCHEDULED', 'UNSCHEDULED')),
        'scheduled': sum(1 for r in results if r['status'] == 'SCHEDULED'),
        'unscheduled': sum(1 for r in results if r['status'] == 'UNSCHEDULED'),
        'duplicates': sum(1 for r in results if r['status'] == 'DUPLICATE'),
        'failed': sum(1 for r in results if r['status'] == 'FAILED'),
    }
    return {'batch_id': batch_id, 'summary': summary, 'results': results}

